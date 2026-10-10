"""Checks that every generated document is complete and safely marked.

Two rules, both enforced on the finished files (not just the in-memory text):
  1. No unfilled {{ }} placeholders anywhere in the document.
  2. The DRAFT / NOT LEGAL ADVICE banner is present on every page.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from src.config import SAFETY_HEADER

# A complete {{ ... }} pair, or a stray {{ or }} left over from a broken template.
UNFILLED = re.compile(r"\{\{.*?\}\}|\{\{|\}\}", re.DOTALL)


class ValidationError(Exception):
    """A generated document failed a safety or completeness check."""


class UnfilledPlaceholderError(ValidationError):
    pass


class MissingSafetyHeaderError(ValidationError):
    pass


def find_unfilled_placeholders(text: str) -> list[str]:
    """Return every unfilled placeholder (or stray brace pair) found in the text."""
    return [" ".join(m.group(0).split()) for m in UNFILLED.finditer(text)]


def assert_no_placeholders(text: str, source: str) -> None:
    found = find_unfilled_placeholders(text)
    if found:
        unique = ", ".join(dict.fromkeys(found))
        raise UnfilledPlaceholderError(f"{source} has {len(found)} unfilled placeholder(s): {unique}")


def _collapse(text: str) -> str:
    return " ".join(text.split())


def read_docx_text(path: Path) -> tuple[str, list[str]]:
    """Return (body text including tables, header text for each section)."""
    doc = Document(str(path))
    body = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            body.extend(cell.text for cell in row.cells)
    headers = ["\n".join(p.text for p in section.header.paragraphs) for section in doc.sections]
    footers = ["\n".join(p.text for p in section.footer.paragraphs) for section in doc.sections]
    return "\n".join(body + footers), headers


def read_pdf_pages(path: Path) -> list[str]:
    return [page.extract_text() or "" for page in PdfReader(str(path)).pages]


def validate_docx(path: Path) -> None:
    body, headers = read_docx_text(path)
    assert_no_placeholders(body + "\n" + "\n".join(headers), path.name)
    if not headers or any(_collapse(SAFETY_HEADER) not in _collapse(h) for h in headers):
        raise MissingSafetyHeaderError(f"{path.name} is missing the safety header in at least one section")


def validate_pdf(path: Path) -> None:
    pages = read_pdf_pages(path)
    assert_no_placeholders("\n".join(pages), path.name)
    unmarked = [i for i, text in enumerate(pages, 1) if _collapse(SAFETY_HEADER) not in _collapse(text)]
    if not pages or unmarked:
        raise MissingSafetyHeaderError(f"{path.name} is missing the safety header on page(s) {unmarked or '(no pages)'}")
