"""Write the run report as .docx and .pdf, in the same house style as the letters.

Landscape US Letter (the clients table is wide), the DRAFT safety banner on every
page, and a footer with the page number. Both files are validated after writing.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import TYPE_CHECKING
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.config import FIRM_NAME, SAFETY_HEADER
from src.document_builder import (
    GRAY, NAVY, NAVY_HEX, PDF_GRAY, PDF_NAVY, PDF_RED, RED, RED_HEX, SERIF,
    _border, _page_number_field, _run, _shade, _spacing,
)
from src.validation import validate_docx, validate_pdf

if TYPE_CHECKING:
    from src.reporter import ReportPart

TITLE = "Client Letters Summary"
FOOTER = f"{FIRM_NAME}  ·  Client letters summary  ·  Page "
HEADER_FILL = "E8ECF4"
# Emoji used in the Markdown report; the serif PDF/Word fonts can't draw them.
EMOJI = re.compile("[✅❌⚠️]\\s*")


def _segments(text: str) -> list[tuple[str, bool]]:
    """Split Markdown-ish text into (text, bold) runs, dropping `code` ticks and emoji."""
    text = EMOJI.sub("", text).replace("`", "")
    return [(seg, i % 2 == 1) for i, seg in enumerate(text.split("**")) if seg]


def _plain(text: str) -> str:
    return "".join(seg for seg, _ in _segments(text))


# --------------------------------------------------------------------------- DOCX

def write_report_docx(parts: list[ReportPart], path: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.left_margin = section.right_margin = Inches(0.75)
    section.top_margin, section.bottom_margin = Inches(0.9), Inches(0.8)
    section.header_distance, section.footer_distance = Inches(0.35), Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(10.5)
    for attr in ("w:eastAsia", "w:cs"):
        normal.element.rPr.rFonts.set(qn(attr), SERIF)
    normal.paragraph_format.space_after = Pt(6)

    banner = section.header.paragraphs[0]
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _border(banner, ("top", "left", "bottom", "right"), RED_HEX, size=6, space=3)
    _shade(banner, "FDF2F2")
    _run(banner, SAFETY_HEADER, size=9, bold=True, color=RED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, FOOTER, size=8.5, color=GRAY)
    _page_number_field(footer, 8.5, GRAY)

    for part in parts:
        _docx_part(doc, part)

    props = doc.core_properties
    props.title, props.author, props.comments = TITLE, FIRM_NAME, SAFETY_HEADER
    doc.save(str(path))
    validate_docx(path)
    return path


def _add_text(paragraph, text: str, **kwargs) -> None:
    for seg, bold in _segments(text):
        _run(paragraph, seg, bold=bold or kwargs.get("bold", False), size=kwargs.get("size"), color=kwargs.get("color"))


def _docx_part(doc, part: ReportPart) -> None:
    kind = part.kind
    if kind == "title":
        p = doc.add_paragraph()
        _spacing(p, after=8)
        _border(p, ("bottom",), NAVY_HEX, size=8, space=4)
        _run(p, part.content, size=17, bold=True, color=NAVY)
    elif kind == "heading":
        p = doc.add_paragraph()
        _spacing(p, before=12, after=4, keep_with_next=True)
        _add_text(p, part.content, size=13, bold=True, color=NAVY)
    elif kind == "note":
        p = doc.add_paragraph()
        _add_text(p, part.content, size=9.5, color=GRAY)
    elif kind == "paragraph":
        _add_text(doc.add_paragraph(), part.content)
    elif kind == "bullets":
        for i, item in enumerate(part.content):
            p = doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Inches(0.4), Inches(-0.2)
            _spacing(p, after=6 if i == len(part.content) - 1 else 1)
            _add_text(p, "•  " + item)
    elif kind in ("kv", "table"):
        header, rows = part.header, part.content
        table = doc.add_table(rows=0, cols=len(rows[0]))
        table.style = "Table Grid"
        if header:
            cells = table.add_row().cells
            for cell, text in zip(cells, header):
                _shade(cell.paragraphs[0], HEADER_FILL)
                _run(cell.paragraphs[0], text, bold=True, size=9.5, color=NAVY)
        for row in rows:
            cells = table.add_row().cells
            for j, (cell, text) in enumerate(zip(cells, row)):
                _add_text(cell.paragraphs[0], text, size=9.5, bold=kind == "kv" and j == 0)
        total = 9.5 * 72  # 11in page minus 0.75in margins, in points
        widths = [130, total - 130] if kind == "kv" else _col_widths(header, rows, total)
        table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Pt(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Pt(width)
                _spacing(cell.paragraphs[0], before=1, after=1)
        doc.add_paragraph()
    else:
        raise ValueError(f"Unknown report part '{kind}'")


# --------------------------------------------------------------------------- PDF

_BODY = ParagraphStyle("rbody", fontName="Times-Roman", fontSize=10.5, leading=13.5, spaceAfter=6)
STYLES = {
    "body": _BODY,
    "title": ParagraphStyle("rtitle", parent=_BODY, fontName="Times-Bold", fontSize=17, leading=21, textColor=PDF_NAVY, spaceAfter=4),
    "heading": ParagraphStyle("rheading", parent=_BODY, fontName="Times-Bold", fontSize=13, leading=16, textColor=PDF_NAVY, spaceBefore=12, spaceAfter=5),
    "note": ParagraphStyle("rnote", parent=_BODY, fontSize=9.5, leading=12, textColor=PDF_GRAY),
    "bullet": ParagraphStyle("rbullet", parent=_BODY, leftIndent=24, bulletIndent=10, spaceAfter=2),
    "cell": ParagraphStyle("rcell", parent=_BODY, fontSize=9.5, leading=11.5, spaceAfter=0),
    "cell_head": ParagraphStyle("rcellhead", parent=_BODY, fontName="Times-Bold", fontSize=9.5, leading=11.5, spaceAfter=0, textColor=PDF_NAVY),
}


def _markup(text: str) -> str:
    return "".join(f"<b>{escape(seg)}</b>" if bold else escape(seg) for seg, bold in _segments(text))


def _p(text: str, style: str, **kwargs) -> Paragraph:
    return Paragraph(_markup(text), STYLES[style], **kwargs)


def _col_widths(header: list[str] | None, rows: list[list[str]], total: float) -> list[float]:
    """Share the page width by how long each column's text is, but never split a word (e.g. a case ID)."""
    lines = ([header] if header else []) + rows
    pad = 14  # cell padding plus a little air
    min_w = [max(stringWidth(word, "Times-Bold", 9.5) for r in lines for word in _plain(r[j]).split() or [""]) + pad
             for j in range(len(lines[0]))]
    weights = [min(max(len(_plain(r[j])) for r in lines), 45) + 4 for j in range(len(lines[0]))]
    spare = total - sum(min_w)
    return [m + spare * w / sum(weights) for m, w in zip(min_w, weights)]


def write_report_pdf(parts: list[ReportPart], path: Path) -> Path:
    pagesize = landscape(LETTER)
    margin = 0.75 * inch
    width = pagesize[0] - 2 * margin

    def decorate(canvas, doc) -> None:
        w, h = pagesize
        canvas.saveState()
        canvas.setFont("Times-Bold", 9)
        box_w = canvas.stringWidth(SAFETY_HEADER, "Times-Bold", 9) + 28
        box_x, box_y = (w - box_w) / 2, h - 0.55 * inch
        canvas.setFillColor(colors.HexColor("#FDF2F2"))
        canvas.setStrokeColor(PDF_RED)
        canvas.setLineWidth(0.8)
        canvas.rect(box_x, box_y, box_w, 16, stroke=1, fill=1)
        canvas.setFillColor(PDF_RED)
        canvas.drawCentredString(w / 2, box_y + 5, SAFETY_HEADER)
        canvas.setFont("Times-Roman", 8.5)
        canvas.setFillColor(PDF_GRAY)
        canvas.drawCentredString(w / 2, 0.45 * inch, FOOTER + str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path), pagesize=pagesize, leftMargin=margin, rightMargin=margin,
        topMargin=0.9 * inch, bottomMargin=0.8 * inch,
        title=TITLE, author=FIRM_NAME, keywords=SAFETY_HEADER, creator="legal-document-generator",
    )

    story = []
    for part in parts:
        kind = part.kind
        if kind in ("title", "note"):
            story.append(_p(part.content, kind))
        elif kind in ("heading", "paragraph"):
            # Start a new page rather than leave a heading or label stranded at the bottom.
            story += [CondPageBreak(1.1 * inch), _p(part.content, "body" if kind == "paragraph" else kind)]
        elif kind == "bullets":
            story += [_p(item, "bullet", bulletText="•") for item in part.content] + [Spacer(1, 4)]
        elif kind in ("kv", "table"):
            header, rows = part.header, part.content
            data = ([[_p(h, "cell_head") for h in header]] if header else []) + [
                [_p(f"**{c}**" if kind == "kv" and j == 0 else c, "cell") for j, c in enumerate(row)] for row in rows
            ]
            widths = [1.8 * inch, width - 1.8 * inch] if kind == "kv" else _col_widths(header, rows, width)
            table = Table(data, colWidths=widths, repeatRows=1 if header else 0)
            style = [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C0CC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
            if header:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{HEADER_FILL}")))
            table.setStyle(TableStyle(style))
            story += [table, Spacer(1, 10)]
        else:
            raise ValueError(f"Unknown report part '{kind}'")
    while story and isinstance(story[-1], Spacer):
        story.pop()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    validate_pdf(path)
    return path
