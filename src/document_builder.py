"""Turn a rendered letter (a list of Blocks) into matching .docx and .pdf files.

Both writers share one layout: US Letter, 1-inch side margins, Times serif body,
firm letterhead on page one, the DRAFT safety banner in the header of every page,
and a footer with the case ID and page number.

After writing, both files are re-opened and validated. If either fails, both are
deleted and the error is raised, so a bad document never sits in the outputs folder.
"""

from __future__ import annotations

import logging
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

from src.config import FIRM_ADDRESS, FIRM_CONTACT, FIRM_NAME, SAFETY_HEADER
from src.templates import Block, exhibit_label, letter_text
from src.validation import assert_no_placeholders, validate_docx, validate_pdf

log = logging.getLogger(__name__)

BODY_SIZE = 11.5
NAVY_HEX, RED_HEX, GRAY_HEX = "1F2A44", "9B1C1C", "5F6368"
FOOTER_TEXT = "{firm}  ·  Privileged & Confidential  ·  Case {case_id}  ·  Page "


def build_letter(blocks: list[Block], out_dir: Path, basename: str, *, title: str, case_id: str) -> list[Path]:
    """Write <basename>.docx and <basename>.pdf, validate both, and return their paths."""
    # Fail fast, before anything touches disk.
    assert_no_placeholders(letter_text(blocks), f"{basename} (rendered text)")

    docx_path = out_dir / f"{basename}.docx"
    pdf_path = out_dir / f"{basename}.pdf"
    try:
        write_docx(blocks, docx_path, title=title, case_id=case_id)
        write_pdf(blocks, pdf_path, title=title, case_id=case_id)
        validate_docx(docx_path)
        validate_pdf(pdf_path)
    except Exception:
        for path in (docx_path, pdf_path):
            path.unlink(missing_ok=True)
        raise
    log.info("    Wrote and validated %s.docx and %s.pdf", basename, basename)
    return [docx_path, pdf_path]


# --------------------------------------------------------------------------- DOCX

SERIF = "Times New Roman"
NAVY, RED, GRAY = (RGBColor.from_string(h) for h in (NAVY_HEX, RED_HEX, GRAY_HEX))

# Elements that must come after <w:pBdr> inside <w:pPr>, per the WordprocessingML schema.
_PPR_AFTER_BORDER = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def _border(paragraph, sides: tuple[str, ...], color: str, size: int = 6, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for side in sides:
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), str(size))
        edge.set(qn("w:space"), str(space))
        edge.set(qn("w:color"), color)
        borders.append(edge)
    p_pr.insert_element_before(borders, *_PPR_AFTER_BORDER)


def _shade(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    p_pr.insert_element_before(shading, *_PPR_AFTER_BORDER[1:])


def _run(paragraph, text: str, *, size: float | None = None, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def _page_number_field(paragraph, size: float, color) -> None:
    run = _run(paragraph, "", size=size, color=color)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _spacing(paragraph, before: float = 0, after: float = 0, keep_with_next: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.keep_with_next = keep_with_next


def write_docx(blocks: list[Block], path: Path, *, title: str, case_id: str) -> None:
    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin, section.bottom_margin = Inches(1), Inches(0.9)
    section.header_distance, section.footer_distance = Inches(0.4), Inches(0.4)

    normal = doc.styles["Normal"]
    normal.font.name = SERIF
    normal.font.size = Pt(BODY_SIZE)
    for attr in ("w:eastAsia", "w:cs"):
        normal.element.rPr.rFonts.set(qn(attr), SERIF)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    # Safety banner: in the header, so it repeats on every page.
    banner = section.header.paragraphs[0]
    banner.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _border(banner, ("top", "left", "bottom", "right"), RED_HEX, size=6, space=3)
    _shade(banner, "FDF2F2")
    _run(banner, SAFETY_HEADER, size=9, bold=True, color=RED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(footer, FOOTER_TEXT.format(firm=FIRM_NAME, case_id=case_id), size=8.5, color=GRAY)
    _page_number_field(footer, 8.5, GRAY)

    # Letterhead (first page only, part of the body).
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(p, after=2)
    _run(p, FIRM_NAME.upper(), size=17, bold=True, color=NAVY)
    for i, line in enumerate((FIRM_ADDRESS, FIRM_CONTACT)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(p, after=0 if i == 0 else 20)
        _run(p, line, size=9.5, color=GRAY)
        if i == 1:
            _border(p, ("bottom",), NAVY_HEX, size=8, space=6)

    for block in blocks:
        _write_docx_block(doc, block)

    props = doc.core_properties
    props.title, props.author, props.subject = title, FIRM_NAME, f"Case {case_id}"
    props.comments = SAFETY_HEADER
    doc.save(str(path))


def _write_docx_block(doc, block: Block) -> None:
    kind, content = block.kind, block.content
    lines = content if isinstance(content, list) else [content]

    if kind == "date":
        _spacing(doc.add_paragraph(content), after=14)
    elif kind == "address":
        for i, line in enumerate(lines):
            _spacing(doc.add_paragraph(line), after=14 if i == len(lines) - 1 else 0)
    elif kind == "subject":
        p = doc.add_paragraph()
        _spacing(p, after=12)
        _run(p, content, bold=True)
    elif kind == "heading":
        p = doc.add_paragraph()
        _spacing(p, before=8, after=3, keep_with_next=True)
        _run(p, content, bold=True, color=NAVY)
    elif kind in ("paragraph", "salutation"):
        for para in content.split("\n\n"):
            doc.add_paragraph(para)
    elif kind == "placeholder":
        p = doc.add_paragraph()
        _run(p, content, italic=True).font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif kind in ("numbered", "bullets"):
        for i, item in enumerate(lines, 1):
            p = doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Inches(0.5), Inches(-0.25)
            fmt.tab_stops.add_tab_stop(Inches(0.5))
            _spacing(p, after=10 if i == len(lines) else 3)
            p.add_run(f"{i}.\t{item}" if kind == "numbered" else f"•\t{item}")
    elif kind == "exhibits":
        for i, item in enumerate(lines, 1):
            p = doc.add_paragraph()
            fmt = p.paragraph_format
            fmt.left_indent, fmt.first_line_indent = Inches(1.25), Inches(-1.0)
            fmt.tab_stops.add_tab_stop(Inches(1.25))
            _spacing(p, after=10 if i == len(lines) else 3)
            _run(p, f"Exhibit {exhibit_label(i)}", bold=True)
            p.add_run(f"\t{item}")
    elif kind == "signature":
        for i, line in enumerate(lines):
            p = doc.add_paragraph(line)
            last = i == len(lines) - 1
            _spacing(p, before=8 if line == "" else 0, after=12 if last else 0, keep_with_next=not last)
    elif kind == "sign_line":
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(3.9))
        _spacing(p, before=26, after=0, keep_with_next=True)
        p.add_run("_" * 36 + "\tDate: " + "_" * 16)
        for line in lines:
            _spacing(doc.add_paragraph(line), after=0)
    else:
        raise ValueError(f"Unknown block kind '{kind}'")


# --------------------------------------------------------------------------- PDF

PDF_NAVY, PDF_RED, PDF_GRAY = (colors.HexColor(f"#{h}") for h in (NAVY_HEX, RED_HEX, GRAY_HEX))

_BODY = ParagraphStyle("body", fontName="Times-Roman", fontSize=BODY_SIZE, leading=15.5, spaceAfter=8)
STYLES = {
    "body": _BODY,
    "firm": ParagraphStyle("firm", fontName="Times-Bold", fontSize=17, leading=21, alignment=TA_CENTER, textColor=PDF_NAVY, spaceAfter=2),
    "letterhead": ParagraphStyle("letterhead", fontName="Times-Roman", fontSize=9.5, leading=12, alignment=TA_CENTER, textColor=PDF_GRAY),
    "tight": ParagraphStyle("tight", parent=_BODY, spaceAfter=0),
    "subject": ParagraphStyle("subject", parent=_BODY, fontName="Times-Bold", spaceAfter=12),
    "heading": ParagraphStyle("heading", parent=_BODY, fontName="Times-Bold", textColor=PDF_NAVY, spaceBefore=8, spaceAfter=3, keepWithNext=1),
    "placeholder": ParagraphStyle("placeholder", parent=_BODY, fontName="Times-Italic", backColor=colors.HexColor("#FFF3B0"), borderPadding=(2, 4, 3, 4), spaceBefore=2, spaceAfter=12),
    "item": ParagraphStyle("item", parent=_BODY, leftIndent=36, bulletIndent=18, spaceAfter=3),
    "exhibit": ParagraphStyle("exhibit", parent=_BODY, leftIndent=90, bulletIndent=18, bulletFontName="Times-Bold", spaceAfter=3),
}


def _p(text: str, style: str, **kwargs) -> Paragraph:
    return Paragraph(escape(text), STYLES[style], **kwargs)


def write_pdf(blocks: list[Block], path: Path, *, title: str, case_id: str) -> None:
    def decorate(canvas, doc) -> None:
        width, height = LETTER
        canvas.saveState()
        # Safety banner, drawn on every page.
        canvas.setFont("Times-Bold", 9)
        box_w = canvas.stringWidth(SAFETY_HEADER, "Times-Bold", 9) + 28
        box_x, box_y = (width - box_w) / 2, height - 0.62 * inch
        canvas.setFillColor(colors.HexColor("#FDF2F2"))
        canvas.setStrokeColor(PDF_RED)
        canvas.setLineWidth(0.8)
        canvas.rect(box_x, box_y, box_w, 16, stroke=1, fill=1)
        canvas.setFillColor(PDF_RED)
        canvas.drawCentredString(width / 2, box_y + 5, SAFETY_HEADER)
        # Footer.
        canvas.setFont("Times-Roman", 8.5)
        canvas.setFillColor(PDF_GRAY)
        canvas.drawCentredString(width / 2, 0.5 * inch, FOOTER_TEXT.format(firm=FIRM_NAME, case_id=case_id) + str(doc.page))
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=0.9 * inch,
        title=title,
        author=FIRM_NAME,
        subject=f"Case {case_id}",
        keywords=SAFETY_HEADER,
        creator="legal-document-generator",
    )

    story = [
        _p(FIRM_NAME.upper(), "firm"),
        _p(FIRM_ADDRESS, "letterhead"),
        _p(FIRM_CONTACT, "letterhead"),
        Spacer(1, 6),
        HRFlowable(width="100%", thickness=0.8, color=PDF_NAVY, spaceAfter=18),
    ]
    for block in blocks:
        story.extend(_pdf_block(block))
    # A trailing gap near the bottom of a page would otherwise spill onto a new, empty page.
    while story and isinstance(story[-1], Spacer):
        story.pop()

    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)


def _pdf_block(block: Block) -> list:
    kind, content = block.kind, block.content
    lines = content if isinstance(content, list) else [content]

    if kind == "date":
        return [_p(content, "body"), Spacer(1, 6)]
    if kind == "address":
        return [_p(line, "tight") for line in lines] + [Spacer(1, 14)]
    if kind == "subject":
        return [_p(content, "subject")]
    if kind == "heading":
        return [_p(content, "heading")]
    if kind in ("paragraph", "salutation"):
        return [_p(para, "body") for para in content.split("\n\n")]
    if kind == "placeholder":
        return [_p(content, "placeholder")]
    if kind in ("numbered", "bullets"):
        items = [_p(item, "item", bulletText=f"{i}." if kind == "numbered" else "•") for i, item in enumerate(lines, 1)]
        return items + [Spacer(1, 6)]
    if kind == "exhibits":
        items = [_p(item, "exhibit", bulletText=f"Exhibit {exhibit_label(i)}") for i, item in enumerate(lines, 1)]
        return items + [Spacer(1, 6)]
    if kind == "signature":
        parts = [Spacer(1, 22) if line == "" else _p(line, "tight") for line in lines]
        return [KeepTogether(parts), Spacer(1, 12)]
    if kind == "sign_line":
        rule = "_" * 36 + " " * 12 + "Date: " + "_" * 16
        return [KeepTogether([Spacer(1, 26), _p(rule, "tight")] + [_p(line, "tight") for line in lines])]
    raise ValueError(f"Unknown block kind '{kind}'")
